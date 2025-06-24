import os
from openpyxl import load_workbook

root_dir = "D:/OneDrive/Desktop/RC CONSORCIO 2025"

for carpeta in os.listdir(root_dir):
    path_junio = os.path.join(root_dir, carpeta, "Junio")
    if os.path.isdir(path_junio):
        for archivo in os.listdir(path_junio):
            if archivo.lower().endswith(".xlsx") and "junio" in archivo.lower():
                ruta_excel = os.path.join(path_junio, archivo)
                try:
                    wb = load_workbook(ruta_excel)
                    cambiado = False
                    for hoja in wb.worksheets:
                        for fila in hoja.iter_rows():
                            for celda in fila:
                                if celda.value and isinstance(celda.value, str):
                                    if "vto 16/06" in celda.value.lower():
                                        celda.value = "VTO 14/07"
                                        cambiado = True
                    if cambiado:
                        wb.save(ruta_excel)
                        print(f"✅ Modificado: {ruta_excel}")
                    else:
                        print(f"➡️  Sin cambios: {ruta_excel}")
                except Exception as e:
                    print(f"❌ Error con {ruta_excel}: {e}")
